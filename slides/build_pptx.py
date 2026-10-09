#!/usr/bin/env python3
"""Build the editable roadshow deck. Copy matches slides/index.html."""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

W = Inches(13.333)
H = Inches(7.5)
RED = RGBColor(0xE2, 0x23, 0x1A)
RED_INK = RGBColor(0xB1, 0x0E, 0x0E)
INK = RGBColor(0x1A, 0x1A, 0x1A)
BODY = RGBColor(0x33, 0x33, 0x33)
MUTED = RGBColor(0x4D, 0x4D, 0x4D)
BLUE = RGBColor(0x0C, 0x5F, 0x86)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
CARD = RGBColor(0xF4, 0xF5, 0xF6)
WASH = RGBColor(0xFF, 0xF6, 0xF5)
LINE = RGBColor(0xE6, 0xE6, 0xE6)
GRAY = RGBColor(0x5E, 0x65, 0x6D)
NAVY = RGBColor(0x16, 0x4E, 0x72)
DARK = RGBColor(0x1C, 0x1C, 0x1C)
SANS = "Source Sans 3"
SERIF = "Source Serif 4"


def set_run(run, text, size, font, color, bold=False):
    run.text = text
    run.font.size = Pt(size)
    run.font.name = font
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.italic = False


def add_text(slide, l, t, w, h, lines, anchor=MSO_ANCHOR.TOP):
    shape = slide.shapes.add_textbox(l, t, w, h)
    tf = shape.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    tf.margin_left = Emu(0)
    tf.margin_right = Emu(0)
    tf.margin_top = Emu(0)
    tf.margin_bottom = Emu(0)
    tf.paragraphs[0].alignment = PP_ALIGN.LEFT
    shape.text_frame.paragraphs[0].space_before = Pt(0)
    shape.text_frame.paragraphs[0].space_after = Pt(0)
    try:
        tf._txBody.bodyPr.set("anchor", {MSO_ANCHOR.TOP: "t", MSO_ANCHOR.MIDDLE: "ctr", MSO_ANCHOR.BOTTOM: "b"}[anchor])
    except Exception:
        pass
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = line.get("align", PP_ALIGN.LEFT)
        p.space_before = Pt(line.get("before", 0))
        p.space_after = Pt(line.get("after", 0))
        run = p.add_run()
        set_run(run, line["text"], line.get("size", 14), line.get("font", SANS), line.get("color", INK), line.get("bold", False))
    return shape


def rect(slide, l, t, w, h, fill, line=None):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, l, t, w, h)
    shape.shadow.inherit = False
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = line
        shape.line.width = Pt(1)
    # drop the default style line if any
    sp = shape._element
    style = sp.find(qn("p:style"))
    if style is not None:
        sp.remove(style)
    return shape


def footer(slide, num, left=0.55):
    add_text(slide, Inches(left), Inches(7.05), Inches(10.5), Inches(0.32), [
        {"text": "Birlasoft | Cyber Underwriting Intelligence | Confidential", "size": 11, "color": RGBColor(0x3F, 0x3F, 0x3F)}
    ])
    add_text(slide, Inches(12.2), Inches(7.05), Inches(0.7), Inches(0.32), [
        {"text": num, "size": 12, "bold": True, "color": INK, "align": PP_ALIGN.RIGHT}
    ])


def rule(slide):
    rect(slide, 0, 0, W, Inches(0.06), RED)


def logo(slide):
    path = Path(__file__).resolve().parents[1] / "birlasoft-logo.png"
    slide.shapes.add_picture(str(path), Inches(11.55), Inches(0.18), height=Inches(0.38))


def eyebrow(slide, text, l=Inches(0.55), t=Inches(0.28)):
    add_text(slide, l, t, Inches(10), Inches(0.28), [
        {"text": text.upper(), "size": 11, "bold": True, "color": RED_INK, "font": SANS}
    ])


def title(slide, text, t=Inches(0.52), size=28, w=Inches(11)):
    add_text(slide, Inches(0.55), t, w, Inches(0.85), [
        {"text": text, "size": size, "font": SERIF, "color": INK}
    ])


def chain(slide, active):
    labels = [
        ("01", "Intake", "Document portal"),
        ("02", "Enrich", "External signals"),
        ("03", "Analyze", "Risk rules"),
        ("04", "Sync", "Policy admin"),
        ("05", "Decide", "Human sign-off"),
    ]
    x = Inches(0.55)
    y = Inches(2.15)
    bw = Inches(2.35)
    gap = Inches(0.12)
    for i, (n, name, desc) in enumerate(labels):
        on = i == active
        rect(slide, x, y, bw, Inches(0.95), WASH if on else RGBColor(0xFA, 0xFA, 0xFA), RED if on else LINE)
        tag = f"{n}  ·  NOW" if on else n
        add_text(slide, x + Inches(0.12), y + Inches(0.08), bw - Inches(0.2), Inches(0.8), [
            {"text": tag, "size": 10, "bold": True, "color": RED_INK if on else RGBColor(0x66, 0x66, 0x66)},
            {"text": name, "size": 16, "bold": True, "color": INK, "before": 2},
            {"text": desc, "size": 12, "color": BODY, "before": 1},
        ])
        x += bw + gap


def points(slide, items):
    x = Inches(0.55)
    y = Inches(3.28)
    bw = Inches(4.0)
    gap = Inches(0.14)
    for title_text, body in items:
        rect(slide, x, y, bw, Inches(2.55), WHITE, RGBColor(0xF0, 0xD0, 0xCC))
        add_text(slide, x + Inches(0.18), y + Inches(0.16), bw - Inches(0.36), Inches(2.25), [
            {"text": title_text, "size": 16, "bold": True, "color": INK},
            {"text": body, "size": 13, "color": BODY, "before": 6},
        ])
        x += bw + gap


def new(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


def build():
    prs = Presentation()
    prs.slide_width = W
    prs.slide_height = H
    prs.core_properties.title = "Adaptive cyber insurance underwriting intelligence"
    prs.core_properties.subject = "Birlasoft cyber underwriting leadership deck"
    prs.core_properties.category = "Confidential"

    # 1 Executive
    s = new(prs)
    rect(s, 0, 0, W, H, WHITE)
    rule(s)
    logo(s)
    eyebrow(s, "Executive briefing")
    title(s, "Adaptive cyber insurance underwriting intelligence", size=30)
    add_text(s, Inches(0.55), Inches(1.35), Inches(11), Inches(0.35), [
        {"text": "Reimagining cyber risk intelligence", "size": 16, "color": BODY}
    ])
    kpis = [
        (False, "INDUSTRY", "3–5 hours", "Typical manual prep for one complex cyber submission."),
        (False, "INDUSTRY", "Same day", "The broker bar. About half a day slower is often enough to lose the placement."),
        (True, "BIRLASOFT", "Same sitting", "A complete package is classified and extracted into Policy Documents, without a rekey bench."),
        (True, "BIRLASOFT", "Human before bind", "Quote, escalate, and decline stay blocked until a person signs off, then sync to policy admin."),
    ]
    x = Inches(0.55)
    for brand, tag, head, body in kpis:
        rect(s, x, Inches(1.85), Inches(3.0), Inches(0.08), RED if brand else RGBColor(0x8D, 0x8D, 0x8D))
        rect(s, x, Inches(1.93), Inches(3.0), Inches(1.7), WASH if brand else CARD)
        add_text(s, x + Inches(0.14), Inches(2.05), Inches(2.72), Inches(1.5), [
            {"text": tag, "size": 10, "bold": True, "color": RED_INK if brand else RGBColor(0x3F, 0x3F, 0x3F)},
            {"text": head, "size": 20, "font": SERIF, "color": INK, "before": 4},
            {"text": body, "size": 12, "color": BODY, "before": 4},
        ])
        x += Inches(3.12)
    add_text(s, Inches(0.55), Inches(3.75), Inches(12), Inches(0.32), [
        {"text": "Industry figures are market observations. Birlasoft figures are workbench design targets, not audited client results.", "size": 11, "color": MUTED}
    ])
    eyebrow(s, "What’s included", t=Inches(4.15))
    features = [
        ("01", "Adaptive questionnaires", "Core, High Limit ($10M+), Renewal Lite, and SME Light."),
        ("02", "Exhibit intelligence", "Classify the application, SOC, IR plan, DR drill, ISO report, and loss runs."),
        ("03", "External control signals", "MFA, EDR, backups, and vendors sit beside the attestation."),
        ("04", "Third-party-compatible rules", "A control gap can block straight-through. An outside score uses the same slot."),
        ("05", "Synchronous policy admin", "One record, a human sign-off, then the handoff to policy admin."),
    ]
    x = Inches(0.55)
    for n, head, body in features:
        rect(s, x, Inches(4.5), Inches(0.06), Inches(2.15), RED)
        rect(s, x + Inches(0.06), Inches(4.5), Inches(2.34), Inches(2.15), WHITE, RGBColor(0xEA, 0xD7, 0xD5))
        add_text(s, x + Inches(0.16), Inches(4.6), Inches(2.16), Inches(1.95), [
            {"text": n, "size": 11, "bold": True, "color": RED_INK},
            {"text": head, "size": 13, "bold": True, "color": INK, "before": 3},
            {"text": body, "size": 11, "color": BODY, "before": 3},
        ])
        x += Inches(2.5)
    footer(s, "01")

    # 2 Barriers
    s = new(prs)
    rect(s, 0, 0, W, H, WHITE)
    rect(s, 0, 0, Inches(4.7), H, RGBColor(0x16, 0x18, 0x1C))
    add_text(s, Inches(0.4), Inches(0.55), Inches(3.9), Inches(3.2), [
        {"text": "THE QUOTE PROBLEM", "size": 12, "bold": True, "color": RGBColor(0xFF, 0xB4, 0xAE)},
        {"text": "Cyber files still arrive as email", "size": 28, "font": SERIF, "color": WHITE, "before": 10},
        {"text": "Every team quoting cyber hits the same four walls — and each one is cleared by hand.", "size": 14, "color": RGBColor(0xEC, 0xEC, 0xEC), "before": 10},
    ])
    rect(s, Inches(0.4), Inches(5.15), Inches(3.7), Inches(1.35), RED)
    add_text(s, Inches(0.55), Inches(5.3), Inches(3.4), Inches(1.1), [
        {"text": "Every submission", "size": 20, "font": SERIF, "color": WHITE},
        {"text": "touched manually before a quote", "size": 14, "color": WHITE, "before": 4},
    ])
    logo(s)
    add_text(s, Inches(5.05), Inches(0.42), Inches(7.5), Inches(1.15), [
        {"text": "CONTEXT", "size": 12, "bold": True, "color": RED_INK},
        {"text": "What blocks a cyber quote today", "size": 26, "font": SERIF, "color": INK, "before": 6},
    ])
    barriers = [
        ("01", "Unstructured packages", "Apps, SOC reports, IR plans, architecture diagrams, and loss runs arrive as mixed email attachments."),
        ("02", "Questionnaires that move", "MFA, EDR, backup, privacy, and vendor questions differ by appetite and go stale as threats change."),
        ("03", "Evidence gaps stall the quote", "A missing restore test or VPN MFA becomes a multi-day broker chase."),
        ("04", "No operational record", "Status lives in inboxes, so quote, referral, and decline never reach policy admin with an audit trail."),
    ]
    y = Inches(1.75)
    for n, head, body in barriers:
        rect(s, Inches(5.05), y, Inches(0.08), Inches(1.15), RED)
        rect(s, Inches(5.13), y, Inches(7.55), Inches(1.15), CARD)
        add_text(s, Inches(5.3), y + Inches(0.14), Inches(0.6), Inches(0.4), [
            {"text": n, "size": 16, "bold": True, "color": RED_INK}
        ])
        add_text(s, Inches(5.95), y + Inches(0.12), Inches(6.5), Inches(0.95), [
            {"text": head, "size": 15, "bold": True, "color": INK},
            {"text": body, "size": 12, "color": BODY, "before": 2},
        ])
        y += Inches(1.25)
    footer(s, "02", 5.05)

    # 3 Market
    s = new(prs)
    rect(s, 0, 0, W, H, WHITE)
    rule(s)
    logo(s)
    eyebrow(s, "Market opportunity  ·  MGA / MGU")
    title(s, "Four observations the market already shares", size=28)
    add_text(s, Inches(0.55), Inches(1.4), Inches(11), Inches(0.35), [
        {"text": "Where a managing general agent can take the quote without growing the rekey bench.", "size": 15, "color": BODY}
    ])
    cards = [
        ("01", "Speed is the product", "In-appetite cyber is won by the market that answers first.", "The placement goes to the quote that arrives while the controls are still current."),
        ("02", "Volume without a larger bench", "Submission volume can grow without a matching underwriting-assistant team.", "Compute takes the reading. People keep the judgment."),
        ("03", "Signals replace the chase", "External signals fill control gaps that used to wait on the insured’s IT team.", "MFA, EDR, and backup evidence can be checked when the file arrives."),
        ("04", "One record for authority", "Delegated authority needs one validated record from intake through quote and policy-admin sync.", "Authority and the audit fail if the handoff is a second rekey."),
    ]
    positions = [(0.55, 1.95), (6.85, 1.95), (0.55, 4.4), (6.85, 4.4)]
    for (l, t), (n, head, body, out) in zip(positions, cards):
        rect(s, Inches(l), Inches(t), Inches(6.0), Inches(2.25), WHITE, RGBColor(0xF0, 0xC9, 0xC4))
        add_text(s, Inches(l + 0.22), Inches(t + 0.14), Inches(5.55), Inches(1.95), [
            {"text": n, "size": 11, "bold": True, "color": RED_INK},
            {"text": head, "size": 18, "bold": True, "color": INK, "before": 4},
            {"text": body, "size": 13, "color": BODY, "before": 4},
            {"text": out, "size": 13, "bold": True, "color": BLUE, "before": 8},
        ])
    footer(s, "03")

    solutions = [
        ("Document intake portal", "Step 1 of 5", 0,
         "Broker mail and every attachment land on the workbench. Nothing is queued for a person to retype.",
         [
             ("Land the package", "Upload or broker mail opens a submission. The intake agent starts from the exhibits, not from a blank form."),
             ("Classify every exhibit", "Application, architecture, SOC report, incident response, disaster recovery, ISO, compliance packet, and loss runs."),
             ("Extract the record", "Named insured, broker, sector, limit, and control fields land in Policy Documents. Adaptive modules — Core, High Limit ($10M+), Renewal Lite, SME Light — preview the broker path."),
         ],
         "Ops owns completeness and the chase until Documents Complete, then hands the file to the underwriter.",
         "04"),
        ("External risk enrichment", "Step 2 of 5", 1,
         "Control signals sit beside the attestation, so a gap is visible before the underwriter opens the PDF.",
         [
             ("Signals beside the form", "MFA, EDR, immutable backups, and critical vendors are read next to what the broker attested."),
             ("Gaps stay on the file", "A remote-access MFA miss stays visible. It does not disappear into an attachment."),
             ("The book, not only the file", "Shared managed-service providers and shared identity providers are flagged for the portfolio."),
         ],
         "Facts that used to wait on the insured’s IT team are on the submission when it is reviewed.",
         "05"),
        ("Risk analysis, third-party compatible", "Step 3 of 5", 2,
         "Triage rules read the ingested file and can block straight-through until a person owns the judgment.",
         [
             ("Rules from the documents", "Triage runs on what was ingested. An open MFA gap blocks a straight-through quote."),
             ("A judgment with a name", "Ignore, block, and escalate are explicit. Each one keeps a rationale on the file."),
             ("Room for an outside score", "The same rule slot accepts a third-party score or vendor model. The MGA is not locked to one security-rating provider."),
         ],
         "Identical risks follow the same encoded rules, whoever supplied the score.",
         "06"),
        ("Policy admin, synchronous", "Step 4 of 5", 3,
         "Limit, retention, and premium are confirmed on the same submission that was ingested.",
         [
             ("Authority before the quote", "Aggregate limit, self-insured retention, and premium are confirmed under financial authority."),
             ("Decisions stay locked", "Quote, escalate, and decline open only after that sign-off is recorded."),
             ("One handoff", "A decline can be undone before PAS sync. Policy admin receives the structured submission, not a second rekey."),
         ],
         "Delegated authority and the admin system share one validated record. A live PAS API is the next step.",
         "07"),
    ]
    for head, step, active, lede, cards, outcome, num in solutions:
        s = new(prs)
        rect(s, 0, 0, W, H, WHITE)
        rule(s)
        logo(s)
        eyebrow(s, f"The solution  ·  Agentic chain  ·  {step}")
        title(s, head, size=28)
        add_text(s, Inches(0.55), Inches(1.45), Inches(12), Inches(0.5), [
            {"text": lede, "size": 15, "color": BODY}
        ])
        chain(s, active)
        points(s, cards)
        add_text(s, Inches(0.55), Inches(6.0), Inches(12), Inches(0.7), [
            {"text": outcome, "size": 14, "bold": True, "color": BLUE}
        ])
        footer(s, num)

    # 8 Human loops
    s = new(prs)
    rect(s, 0, 0, W, H, WHITE)
    rule(s)
    logo(s)
    eyebrow(s, "The solution  ·  Agentic chain  ·  Step 5 of 5")
    title(s, "Human validation loops", size=28)
    add_text(s, Inches(0.55), Inches(1.45), Inches(12), Inches(0.4), [
        {"text": "Agents prepare the file. People own completeness, risk, and the decision to quote.", "size": 15, "color": BODY}
    ])
    chain(s, 4)
    loops = [
        ("OPS LOOP", "Make the package whole", ["Package completeness", "Missing-document chase", "Broker reminder", "Ready for UW"]),
        ("DECISION DESK", "Own the risk decision", ["Risk narrative", "Financial authority", "Quote, escalate, or decline", "Specialist referral"]),
    ]
    x = Inches(0.55)
    for label, head, items in loops:
        rect(s, x, Inches(3.28), Inches(6.05), Inches(2.7), CARD)
        lines = [
            {"text": label, "size": 11, "bold": True, "color": RED_INK},
            {"text": head, "size": 18, "bold": True, "color": INK, "before": 4},
        ]
        for item in items:
            lines.append({"text": "–  " + item, "size": 14, "color": INK, "before": 4})
        add_text(s, x + Inches(0.22), Inches(3.42), Inches(5.6), Inches(2.45), lines)
        x += Inches(6.2)
    add_text(s, Inches(0.55), Inches(6.15), Inches(12), Inches(0.55), [
        {"text": "Tier assist is a recommendation. It is not a binding premium. Every path keeps a rationale.", "size": 14, "bold": True, "color": BLUE}
    ])
    footer(s, "08")

    # 9 Architecture
    s = new(prs)
    rect(s, 0, 0, W, H, WHITE)
    rule(s)
    logo(s)
    eyebrow(s, "Solution overview")
    title(s, "Architecture and roadmap checklist", size=26, t=Inches(0.42))
    add_text(s, Inches(0.55), Inches(1.15), Inches(12), Inches(0.3), [
        {"text": "Existing channels stay. The chain connects intake, the underwriter, and policy admin.", "size": 13, "color": BODY}
    ])
    cols = [
        (GRAY, "Broker channels", ["Broker email", "Upload on the workbench", "Web and APIs"]),
        (RED, "Agentic chain", ["Intake and classification", "External enrichment", "Risk rules", "Human sign-off"]),
        (NAVY, "Policy admin", ["Limit, SIR, premium", "Financial authority", "Quote, escalate, decline", "PAS-style sync"]),
        (GRAY, "External ecosystem", ["Security-rating vendors", "Brokers and MGAs", "Carriers and capacity"]),
    ]
    x = Inches(0.45)
    for i, (color, head, items) in enumerate(cols):
        bw = Inches(2.85)
        rect(s, x, Inches(1.52), bw, Inches(0.38), color)
        add_text(s, x, Inches(1.56), bw, Inches(0.32), [
            {"text": head, "size": 12, "bold": True, "color": WHITE, "align": PP_ALIGN.CENTER}
        ])
        rect(s, x, Inches(1.9), bw, Inches(1.35), WASH if color == RED else CARD, RGBColor(0xF0, 0xC2, 0xBE) if color == RED else None)
        add_text(s, x + Inches(0.12), Inches(1.96), bw - Inches(0.2), Inches(1.2), [
            {"text": item, "size": 11, "color": INK, "before": 1 if n else 0} for n, item in enumerate(items)
        ])
        if i < 3:
            rect(s, x + bw + Inches(0.05), Inches(2.4), Inches(0.28), Inches(0.04), RED)
        x += Inches(3.18)
    for i, label in enumerate(["Access control", "Audit and change logs", "Dashboards and reporting"]):
        rect(s, Inches(0.45 + i * 4.25), Inches(3.38), Inches(4.05), Inches(0.32), DARK)
        add_text(s, Inches(0.45 + i * 4.25), Inches(3.4), Inches(4.05), Inches(0.28), [
            {"text": label, "size": 11, "bold": True, "color": WHITE, "align": PP_ALIGN.CENTER}
        ])
    add_text(s, Inches(0.45), Inches(3.74), Inches(12.4), Inches(0.36), [
        {"text": "In the workbench today: intake, extraction, enrichment, risk rules, human sign-off, and PAS-style handoff. Live rating-vendor connectors and a production PAS API are the next expansion.", "size": 10, "color": MUTED}
    ])
    rows = [
        ["Capability", "Manual inbox", "Point extraction tool", "Birlasoft", "Horizon"],
        ["Mixed-exhibit intake", "Rekey from email", "Parse a single file", "Classify and extract the package", "Now"],
        ["Adaptive questionnaires", "One static form", "Fixed template", "Core, High Limit, Renewal Lite, SME Light", "Now"],
        ["Control-signal enrichment", "Chase by email", "Paste in a score", "Signals beside the attestation", "Now"],
        ["Vendor-neutral risk rules", "Underwriter memory", "One locked model", "Rule slot accepts an external score", "Now"],
        ["Human sign-off before bind", "Informal review", "Often skipped", "Quote, escalate, and decline stay locked", "Now"],
        ["Synchronous PAS handoff", "Second rekey", "File export", "Same record now; live API next", "Next"],
        ["Portfolio view", "Spreadsheet", "Single file only", "Shared MSP and identity-provider watch", "Now"],
    ]
    table = s.shapes.add_table(len(rows), 5, Inches(0.4), Inches(4.12), Inches(12.5), Inches(2.75)).table
    widths = [Inches(2.35), Inches(2.15), Inches(2.25), Inches(4.35), Inches(1.4)]
    for i, width in enumerate(widths):
        table.columns[i].width = width
    for r, row in enumerate(rows):
        for c, val in enumerate(row):
            cell = table.cell(r, c)
            cell.text = ""
            cell.margin_left = Inches(0.06)
            cell.margin_right = Inches(0.04)
            cell.margin_top = Inches(0.04)
            cell.margin_bottom = Inches(0.02)
            p = cell.text_frame.paragraphs[0]
            run = p.add_run()
            bold = r == 0 or c == 4 or c == 3
            color = RED_INK if c == 4 and r else (INK if r == 0 else INK)
            set_run(run, val, 10 if r else 9, SANS, color, bold=bold)
            fill = WASH if c == 3 else WHITE
            if r == 0:
                fill = RGBColor(0xF7, 0xF7, 0xF7)
            cell.fill.solid()
            cell.fill.fore_color.rgb = fill
    footer(s, "09")

    # 10 Close
    s = new(prs)
    rect(s, 0, 0, W, H, WHITE)
    rule(s)
    logo(s)
    eyebrow(s, "Conclusion and next steps")
    title(s, "More time for underwriting", size=32)
    add_text(s, Inches(0.55), Inches(1.5), Inches(12), Inches(0.45), [
        {"text": "Agentic intake, enrichment, and validation so underwriters price cyber risk instead of chasing documents.", "size": 15, "color": BODY}
    ])
    ctas = [
        ("01", "Talk to a Birlasoft advisor", "Share your cyber intake challenges and identify a focused workflow to explore."),
        ("02", "See a submission in the workbench", "Walk a broker package through to a structured record.\nOpen the live demo: birlasoftdemo.github.io/cyber-uw"),
        ("03", "Agree how to measure value", "Define success criteria together and evaluate the outcomes before scaling."),
    ]
    x = Inches(0.55)
    for n, head, body in ctas:
        rect(s, x, Inches(2.15), Inches(4.0), Inches(2.35), CARD)
        add_text(s, x + Inches(0.2), Inches(2.3), Inches(3.6), Inches(2.05), [
            {"text": n, "size": 14, "bold": True, "color": RED_INK},
            {"text": head, "size": 16, "bold": True, "color": INK, "before": 6},
            {"text": body, "size": 13, "color": BODY, "before": 6},
        ])
        x += Inches(4.15)
    add_text(s, Inches(0.55), Inches(4.8), Inches(12), Inches(1.2), [
        {"text": "Contact us to plan your next step", "size": 24, "font": SERIF, "color": BLUE},
        {"text": "Reach out to the Birlasoft team to discuss your cyber underwriting needs.", "size": 15, "color": BODY, "before": 6},
    ])
    footer(s, "10")

    out = str(Path(__file__).resolve().parent / "cyber-uw-intelligence.pptx")
    prs.save(out)
    print(out, len(prs.slides))


if __name__ == "__main__":
    build()
